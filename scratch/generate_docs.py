import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import os

def create_documentation():
    doc = docx.Document()

    # Set Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Color Palette Constants
    HEX_PRIMARY = "1E3A8A"      # Deep Navy Blue
    HEX_SECONDARY = "7C5CFF"    # Electric Purple / Indigo
    HEX_DARK = "111827"         # Dark Charcoal Body Text
    HEX_LIGHT_BG = "F8FAFC"     # Light Slate / Gray Background
    HEX_BORDER = "CBD5E1"       # Soft Slate Border
    HEX_CALLOUT_BG = "F0F4FF"   # Light Blue Tint Callout
    
    RGB_PRIMARY = RGBColor(0x1E, 0x3A, 0x8A)
    RGB_SECONDARY = RGBColor(0x7C, 0x5C, 0xFF)
    RGB_DARK = RGBColor(0x11, 0x18, 0x27)
    RGB_MUTED = RGBColor(0x4B, 0x55, 0x63)
    RGB_WHITE = RGBColor(0xFF, 0xFF, 0xFF)

    # Helper XML functions
    def set_cell_background(cell, hex_color):
        shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading_elm)

    def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'''
            <w:tcMar {nsdecls("w")}>
                <w:top w:w="{top}" w:type="dxa"/>
                <w:bottom w:w="{bottom}" w:type="dxa"/>
                <w:left w:w="{left}" w:type="dxa"/>
                <w:right w:w="{right}" w:type="dxa"/>
            </w:tcMar>
        ''')
        tcPr.append(tcMar)

    def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideV w:val="none"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr.append(borders)

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = RGB_PRIMARY
        
        # Bottom accent border line under Heading 1
        pBdr = parse_xml(f'''
            <w:pBdr {nsdecls("w")}>
                <w:bottom w:val="single" w:sz="12" w:space="4" w:color="{HEX_PRIMARY}"/>
            </w:pBdr>
        ''')
        p._p.get_or_add_pPr().append(pBdr)
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGB_SECONDARY
        return p

    def add_heading_3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGB_DARK
        return p

    def add_body_p(text, bold_prefix="", space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            run_b.font.name = "Arial"
            run_b.font.size = Pt(10.5)
            run_b.font.bold = True
            run_b.font.color.rgb = RGB_DARK
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGB_DARK
        return p

    def add_bullet_p(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            run_b.font.name = "Arial"
            run_b.font.size = Pt(10.5)
            run_b.font.bold = True
            run_b.font.color.rgb = RGB_DARK
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGB_DARK
        return p

    def add_callout(text, title="KEY ARCHITECTURAL INSIGHT", color_hex=HEX_PRIMARY, bg_hex=HEX_CALLOUT_BG):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, bg_hex)
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
        
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="{color_hex}"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        run_t = p.add_run(f"📌 {title}: ")
        run_t.bold = True
        run_t.font.name = "Arial"
        run_t.font.size = Pt(10.5)
        run_t.font.color.rgb = RGB_PRIMARY
        
        run_b = p.add_run(text)
        run_b.font.name = "Arial"
        run_b.font.size = Pt(10)
        run_b.font.color.rgb = RGB_DARK
        doc.add_paragraph() # Spacing

    def add_code_block(code_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, HEX_LIGHT_BG)
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>
                <w:left w:val="single" w:sz="16" w:space="0" w:color="{HEX_SECONDARY}"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>
                <w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(code_text)
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        
        sp_p = doc.add_paragraph()
        sp_p.paragraph_format.space_before = Pt(0)
        sp_p.paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # COVER / HEADER
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(24)
    title_p.paragraph_format.space_after = Pt(6)
    title_run = title_p.add_run("ExcelDraw")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(32)
    title_run.font.bold = True
    title_run.font.color.rgb = RGB_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(16)
    sub_run = sub_p.add_run("Production-Grade Technical Documentation Manual")
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(16)
    sub_run.font.bold = True
    sub_run.font.color.rgb = RGB_SECONDARY

    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(24)
    meta_run = meta_p.add_run("Architectural Specification | Full-Stack Monorepo | Vector Canvas Engine | Real-Time Synchronization Protocol\nDocument Version: 1.0.0 | Target Runtime: Node.js 20.x, PostgreSQL 14+, Next.js 15")
    meta_run.font.name = "Arial"
    meta_run.font.size = Pt(9.5)
    meta_run.font.italic = True
    meta_run.font.color.rgb = RGB_MUTED

    # Bottom dividing rule
    divider_p = doc.add_paragraph()
    divider_p.paragraph_format.space_after = Pt(16)
    pBdr = parse_xml(f'''
        <w:pBdr {nsdecls("w")}>
            <w:bottom w:val="single" w:sz="18" w:space="1" w:color="{HEX_PRIMARY}"/>
        </w:pBdr>
    ''')
    divider_p._p.get_or_add_pPr().append(pBdr)

    # -------------------------------------------------------------
    # SECTION 1: EXECUTIVE SUMMARY & OBJECTIVES
    # -------------------------------------------------------------
    add_heading_1("1. Executive Summary & Objectives")

    add_body_p(
        "ExcelDraw is an enterprise-ready, high-performance collaborative whiteboard and vector diagramming application designed for distributed teams, engineering architects, and product designers. Structured as a modern, unified monorepo, ExcelDraw empowers multiple concurrent users to interact on an infinite, zoomable vector canvas with real-time multi-user state synchronization, near-zero visual latency, and robust offline-first durability.",
        bold_prefix="Overview: "
    )

    add_heading_2("1.1 Core Problems Solved")
    add_bullet_p(
        "Standard web applications struggle with concurrency conflicts and noticeable input latency when multiple users draw simultaneously. ExcelDraw solves this by isolating high-frequency rendering onto a dedicated HTML5 2D Canvas engine with client-side deterministic state reconstruction, transmitting lightweight incremental delta wire messages over WebSockets rather than full canvas binary frames.",
        bold_prefix="Real-Time Visual Latency & Concurrency: "
    )
    add_bullet_p(
        "Monolithic applications suffer from tight coupling between frontend UI components, REST API backend services, real-time message brokers, and shared database models. ExcelDraw addresses this by leveraging a Turborepo monorepo structure, strictly partitioning concerns into dedicated micro-applications and re-usable typed packages.",
        bold_prefix="Monolithic Micro-Frontend Separation: "
    )
    add_bullet_p(
        "Traditional web whiteboards lose element precision when scaling or panning across large diagrams. ExcelDraw implements an infinite viewport transformation engine with arbitrary matrix scaling (0.1x to 8.0x zoom), custom high-DPI scaling (devicePixelRatio adaptation), custom dark/light theme neon rendering, and spatial hit-testing algorithms.",
        bold_prefix="Infinite Vector Canvas Math & High-DPI Rendering: "
    )
    add_bullet_p(
        "Whiteboard drawings must survive server restarts, connection loss, or browser reloads. ExcelDraw implements durable chat stream persistence in PostgreSQL, allowing room participants to instantly replay room history sequentially upon re-entry.",
        bold_prefix="State Durability & Historical Event Replay: "
    )

    add_heading_2("1.2 Primary Use Cases")
    add_bullet_p("Interactive software architecture sketching, sequence flowcharts, database entity mapping, and UML modeling.", bold_prefix="System Architecture & Flowcharting: ")
    add_bullet_p("Distributed remote engineering teams conducting sprint planning, retrospectives, and visual brainstorm sessions.", bold_prefix="Collaborative Team Whiteboarding: ")
    add_bullet_p("Rapid low-fidelity UI mockups, layout drafting, and visual annotation.", bold_prefix="Product Wireframing & Design Drafting: ")
    add_bullet_p("Interactive virtual blackboards for online tutoring, technical interviews, and live presentations.", bold_prefix="Technical Education & Presentations: ")

    add_callout(
        "ExcelDraw separates concerns into three standalone runtime runtimes: a Next.js 15 frontend application (`excel_front`), an Express.js HTTP backend (`http-backend`), and a native WebSocket real-time server (`websockets`). Shared schema types and database connectors are shared across the monorepo compile-time boundary.",
        title="SYSTEM OBJECTIVE"
    )

    # -------------------------------------------------------------
    # SECTION 2: COMPLETE TECH STACK
    # -------------------------------------------------------------
    add_heading_1("2. Complete Tech Stack")

    add_body_p(
        "The following table details every programming language, framework, library, tool, and execution engine detected across all workspace configuration files (`package.json`, `pnpm-workspace.yaml`, `turbo.json`, `render.yaml`, `schema.prisma`)."
    )

    tech_table = doc.add_table(rows=1, cols=4)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tech_table)

    headers = ["Category", "Technology / Library", "Version / Specifier", "Purpose & Implementation Role"]
    hdr_cells = tech_table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].width = Inches([1.3, 1.6, 1.1, 2.5][i])
        set_cell_background(hdr_cells[i], HEX_PRIMARY)
        set_cell_margins(hdr_cells[i], top=140, bottom=140, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h)
        run.font.name = "Arial"
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGB_WHITE

    tech_data = [
        ("Monorepo Architecture", "Turborepo", "^2.3.3", "High-performance monorepo build system & task orchestrator."),
        ("Monorepo Architecture", "pnpm Workspaces", "v9.0.0", "Strict, disk-efficient workspace package manager (`pnpm-workspace.yaml`)."),
        ("Monorepo Architecture", "TypeScript", "^5.5.4", "Strict static typing across packages and applications."),
        ("Monorepo Architecture", "Prettier", "^3.2.5", "Unified code formatting rule enforcement."),
        ("Frontend Application", "Next.js", "15.5.19 (App Router)", "React framework with Server Components, Turbopack, and client routing."),
        ("Frontend Application", "React & React DOM", "^18.3.1", "UI component library and DOM reconciliation engine."),
        ("Frontend Application", "Tailwind CSS & PostCSS", "^3.4.1", "Utility-first CSS framework with custom dark-mode token extensions."),
        ("Frontend Application", "Framer Motion", "^11.18.1", "Production-grade fluid UI transitions and modal animations."),
        ("Frontend Application", "Zustand", "^5.0.3", "Lightweight, un-opinionated state management (canvas theme store)."),
        ("Frontend Application", "Radix UI Primitives", "Various (^1.1 - ^2.0)", "Accessible headless UI primitives (Dialog, Slider, Dropdown, Tabs)."),
        ("Frontend Application", "Lucide React", "^0.473.0", "Consistent vector icons for canvas tools and dashboard interface."),
        ("Frontend Application", "Sonner", "^1.5.0", "Opinionated toast notification system for action feedback."),
        ("Frontend Application", "Axios & JS-Cookie", "^1.7.9 / ^3.0.5", "HTTP client for REST endpoints & browser cookie storage helper."),
        ("HTTP Backend API", "Express.js", "^4.x", "Fast, minimal web framework for user authentication and room REST APIs."),
        ("HTTP Backend API", "JSON Web Token", "^9.x (`jsonwebtoken`)", "Stateless authentication token generation and validation."),
        ("HTTP Backend API", "bcrypt", "^5.x", "Secure password hashing with 10 salt rounds."),
        ("HTTP Backend API", "CORS & Cookie", "Standard Node packages", "Cross-Origin Resource Sharing guard & cookie header parser."),
        ("Real-Time Engine", "ws (Native WebSockets)", "^8.x", "Low-latency WebSocket server handling real-time socket connections."),
        ("Database Layer", "PostgreSQL", "v14+", "Relational database server hosting users, rooms, and chat message logs."),
        ("Database Layer", "Prisma ORM", "^6.2.1 (`@prisma/client`)", "Type-safe database ORM and migration generator."),
        ("Validation & Schema", "Zod", "^3.23.8", "TypeScript-first schema validation for API request bodies."),
        ("Cloud Infrastructure", "Render", "render.yaml Spec", "PaaS deployment configuration for `exceldraw-api` and `exceldraw-ws`."),
        ("Cloud Infrastructure", "Vercel", "vercel.json Spec", "Serverless deployment target for `excel_front` Next.js application.")
    ]

    for idx, row in enumerate(tech_data):
        row_cells = tech_table.add_row().cells
        bg_color = HEX_LIGHT_BG if idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate(row):
            row_cells[i].width = Inches([1.3, 1.6, 1.1, 2.5][i])
            set_cell_background(row_cells[i], bg_color)
            set_cell_margins(row_cells[i], top=100, bottom=100, left=120, right=120)
            p = row_cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(val)
            run.font.name = "Arial"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGB_DARK
            if i == 1:
                run.font.bold = True

    doc.add_paragraph() # Spacing

    # -------------------------------------------------------------
    # SECTION 3: SYSTEM ARCHITECTURE & WORKFLOW
    # -------------------------------------------------------------
    add_heading_1("3. System Architecture & Workflow")

    add_body_p(
        "ExcelDraw is engineered as a decoupled monorepo system consisting of three distinct runtime tiers, backed by shared compile-time package infrastructure. All real-time canvas events are processed asynchronously through a specialized WebSocket stream, while HTTP operations (auth, room administration, initial chat history fetching) are routed through an Express REST API."
    )

    add_heading_2("3.1 Topology & Component Flow Diagram")
    add_body_p("The ASCII structural overview below illustrates the communication pipelines connecting the client browser, API backend, real-time message broker, and database storage layer:")

    arch_diagram = (
        "+-----------------------------------------------------------------------------------+\n"
        "|                                CLIENT BROWSER                                    |\n"
        "|  +--------------------------+  +----------------------+  +---------------------+  |\n"
        "|  |   Next.js 15 UI Layer    |  |  Zustand Theme Store |  | HTML5 2D Canvas Engine|  |\n"
        "|  | (Dashboard, Auth, Room)  |  |   (Dark/Light State) |  |   (Game.ts Class)   |  |\n"
        "|  +-------------+------------+  +----------------------+  +----------+----------+  |\n"
        "+----------------|---------------------------------------------------|--------------+\n"
        "                 |                                                   |\n"
        "       HTTP REST | (Headers: Bearer JWT / Cookie)         WebSockets | (Query: ?token=JWT)\n"
        "                 v                                                   v\n"
        "+----------------------------------+               +--------------------------------+\n"
        "|      EXPRESS HTTP BACKEND        |               |    WEBSOCKET REALTIME SERVER   |\n"
        "|        (apps/http-backend)       |               |        (apps/websockets)        |\n"
        "|  - /signup, /signin (bcrypt Auth)|               |  - Connection Token Verification|\n"
        "|  - /room, /rooms (Slug Admin)    |               |  - User/Room In-Memory Registry |\n"
        "|  - /chats/:roomId (History fetch)|               |  - Realtime Chat/Delta Re-cast  |\n"
        "+----------------+-----------------+               +----------------+---------------+|\n"
        "                 |                                                  |\n"
        "                 | Prisma Client                                    | Prisma Client\n"
        "                 v                                                  v\n"
        "+-----------------------------------------------------------------------------------+\n"
        "|                               POSTGRESQL DATABASE                                 |\n"
        "|          models: User (id, email, password), Room (id, slug), Chat (message)      |\n"
        "+-----------------------------------------------------------------------------------+"
    )
    add_code_block(arch_diagram)

    add_heading_2("3.2 Detailed Inter-Component Workflows")

    add_heading_3("Workflow 1: User Registration & Authentication Phase")
    add_bullet_p("The client submits email, password, firstName, lastName, and optional Avatar to `POST /signup`.", bold_prefix="1. Request Validation: ")
    add_bullet_p("The `http-backend` validates inputs using `CreateUserSchema` (Zod). If invalid, a `400 Bad Request` is returned.", bold_prefix="2. Schema Enforcement: ")
    add_bullet_p("The server hashes the raw password using `bcrypt.hash(password, 10)`. User creation is executed via `@repo/database` Prisma client.", bold_prefix="3. Password Encryption: ")
    add_bullet_p("Upon successful insert, a JWT signed with `JWT_SECRET` (expiry: 7 days) is returned to the client and stored in browser HTTP cookies (`js-cookie`).", bold_prefix="4. Token Issuance: ")

    add_heading_3("Workflow 2: Room Creation & Workspace Initialization")
    add_bullet_p("An authenticated user inputs a unique room name/slug on the `/dashboard` page and submits the form.", bold_prefix="1. User Action: ")
    add_bullet_p("The client sends a `POST /room` request containing `{ name: slug }`. The request is intercepted by `middleware.ts`, which extracts the JWT from the `Authorization: Bearer <token>` header or `cookie.token`.", bold_prefix="2. REST Call & Authorization: ")
    add_bullet_p("The backend verifies the token and creates a record in the `Room` table with `adminId = req.userId`. Returns `201 Created` with the new numeric `roomId`.", bold_prefix="3. Database Provisioning: ")
    add_bullet_p("Next.js router navigates the user to `/canvas/[roomId]`.", bold_prefix="4. Client Redirect: ")

    add_heading_3("Workflow 3: WebSocket Connection & Handshake Guard")
    add_bullet_p("The `/canvas/[roomId]` page mounts the `RoomCanvas` component, which checks for the presence of `token` in cookies.", bold_prefix="1. Guard Verification: ")
    add_bullet_p("`RoomCanvas` instantiates a `WebSocket` connection to `ws://localhost:8080?token=<token>`.", bold_prefix="2. Handshake Initiation: ")
    add_bullet_p("The WebSocket server parses the `token` parameter from `request.url`. If token verification (`jwt.verify`) fails or is missing, the connection is instantly terminated (`ws.close()`).", bold_prefix="3. Server Verification: ")
    add_bullet_p("Upon `ws.onopen`, the client sends `JSON.stringify({ type: 'join_room', roomId: String(roomId) })`. The WS server pushes `{ userId, rooms: [roomId], socket: ws }` into its global `users` memory array.", bold_prefix="4. Room Joining: ")

    add_heading_3("Workflow 4: Real-Time Vector Drawing & Event Synchronization")
    add_bullet_p("When a user draws a rectangle, circle, line, rhombus, or freehand stroke, the client-side `Game` engine constructs a `Shape` object containing a unique `id` (`crypto.randomUUID()`), coordinates, stroke color, and stroke width.", bold_prefix="1. Canvas Event Capture: ")
    add_bullet_p("The `Game` engine invokes `broadcast(shape)`, sending `{ type: 'chat', message: JSON.stringify(shape), roomId }` over the open WebSocket.", bold_prefix="2. Incremental Broadcast: ")
    add_bullet_p("The WebSocket server receives the payload, executes `prismaClient.chat.create(...)` to persist the wire message into the database, and immediately loops through all active `users` joined in `roomId` to send the payload.", bold_prefix="3. Persistence & Re-cast: ")
    add_bullet_p("Remote client sockets receive the `chat` event, parse the nested `WireMessage`, invoke `applyMessage()`, and trigger `game.render()` to redraw the 2D canvas.", bold_prefix="4. Remote Re-rendering: ")

    add_heading_3("Workflow 5: Historical State Replay & Re-hydration")
    add_bullet_p("When a user enters a room, `Game.ts` initializes by executing `getExistingShape(roomId)` (`apps/excel_front/draw/http.ts`).", bold_prefix="1. History Fetch: ")
    add_bullet_p("The client issues a `GET /chats/:roomId` REST request. The server queries the database for up to 10,000 `Chat` rows ordered descending by `id`.", bold_prefix="2. Query Execution: ")
    add_bullet_p("The frontend receives the messages, reverses the array to restore chronological order (oldest to newest), parses each wire message JSON, and feeds them into `applyMessage()` to deterministically build the final canvas state.", bold_prefix="3. State Folding: ")

    # -------------------------------------------------------------
    # SECTION 4: CORE LOGIC & KEY FEATURES
    # -------------------------------------------------------------
    add_heading_1("4. Core Logic & Key Features")

    add_body_p(
        "This section explains the core logic, business rules, data schemas, security measures, and mathematical algorithms driving ExcelDraw."
    )

    add_heading_2("4.1 Database Model & Schema Architecture (`schema.prisma`)")
    add_body_p(
        "The relational schema is maintained in `packages/database/prisma/schema.prisma` using PostgreSQL as the datasource. It consists of three primary entities: `User`, `Room`, and `Chat`."
    )

    db_table = doc.add_table(rows=1, cols=4)
    db_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(db_table)

    db_headers = ["Model Name", "Field & Type", "Attributes / Constraints", "Description & Business Rule"]
    db_hdr_cells = db_table.rows[0].cells
    for i, h in enumerate(db_headers):
        db_hdr_cells[i].width = Inches([1.2, 1.5, 1.5, 2.3][i])
        set_cell_background(db_hdr_cells[i], HEX_PRIMARY)
        set_cell_margins(db_hdr_cells[i], top=140, bottom=140, left=120, right=120)
        p = db_hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h)
        run.font.name = "Arial"
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGB_WHITE

    db_data = [
        ("User", "id String\nemail String\npassword String\nfirstName String\nlastName String\nAvatar String?", "@id @default(cuid())\n@unique\n-\n-\n-\noptional", "Stores user profile credentials. Passwords are stored exclusively as bcrypt hashes. Email must be unique (`P2002` error handling)."),
        ("Room", "id Int\nslug String\ncreatedAt DateTime\nadminId String", "@id @default(autoincrement())\n@unique\n@default(now())\nFK -> User.id", "Represents a drawing workspace room. `slug` is a human-readable unique identifier. `adminId` references the room creator."),
        ("Chat", "id Int\nmessage String\nuserId String\nroomId Int\ncreatedAt DateTime", "@id @default(autoincrement())\nLong text payload\nFK -> User.id\nFK -> Room.id\n@default(now())", "Stores raw JSON wire messages. Holds shape definitions, shape updates, and deletion markers for room persistence and replay.")
    ]

    for idx, row in enumerate(db_data):
        row_cells = db_table.add_row().cells
        bg_color = HEX_LIGHT_BG if idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate(row):
            row_cells[i].width = Inches([1.2, 1.5, 1.5, 2.3][i])
            set_cell_background(row_cells[i], bg_color)
            set_cell_margins(row_cells[i], top=100, bottom=100, left=120, right=120)
            p = row_cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(val)
            run.font.name = "Arial"
            run.font.size = Pt(9.0)
            run.font.color.rgb = RGB_DARK
            if i == 0:
                run.font.bold = True

    doc.add_paragraph()

    add_heading_2("4.2 REST API Specification (`apps/http-backend`)")
    add_body_p("The HTTP API server exposes the following authentication and management endpoints:")

    api_table = doc.add_table(rows=1, cols=4)
    api_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(api_table)

    api_headers = ["HTTP Method", "Endpoint Route", "Auth Required", "Request Body / Description"]
    api_hdr_cells = api_table.rows[0].cells
    for i, h in enumerate(api_headers):
        api_hdr_cells[i].width = Inches([1.1, 1.6, 1.1, 2.7][i])
        set_cell_background(api_hdr_cells[i], HEX_PRIMARY)
        set_cell_margins(api_hdr_cells[i], top=140, bottom=140, left=120, right=120)
        p = api_hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h)
        run.font.name = "Arial"
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGB_WHITE

    api_data = [
        ("POST", "/signup", "No", "Body: `{ email, password, firstName, lastName, Avatar? }`. Validated with `CreateUserSchema`. Returns `{ token, message }`."),
        ("POST", "/signin", "No", "Body: `{ email, password }`. Validated with `SigninSchema`. Verifies bcrypt hash. Returns `{ token }`."),
        ("POST", "/room", "Yes (`middleware`)", "Body: `{ name }`. Validated with `CreateRoomSchema`. Creates room slug for authenticated admin. Returns `{ roomId, message }`."),
        ("GET", "/rooms", "Yes (`middleware`)", "Query: None. Fetches list of all active rooms `{ id, slug, createdAt }` for the dashboard grid."),
        ("GET", "/room/:slug", "Yes (`middleware`)", "Params: `slug`. Resolves room object by unique string slug. Returns `{ room }`."),
        ("GET", "/chats/:roomId", "Yes (`middleware`)", "Params: `roomId`. Returns up to 10,000 recorded wire message objects for room canvas re-hydration.")
    ]

    for idx, row in enumerate(api_data):
        row_cells = api_table.add_row().cells
        bg_color = HEX_LIGHT_BG if idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate(row):
            row_cells[i].width = Inches([1.1, 1.6, 1.1, 2.7][i])
            set_cell_background(row_cells[i], bg_color)
            set_cell_margins(row_cells[i], top=100, bottom=100, left=120, right=120)
            p = row_cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(val)
            run.font.name = "Arial"
            run.font.size = Pt(9.0)
            run.font.color.rgb = RGB_DARK
            if i == 0:
                run.font.bold = True

    doc.add_paragraph()

    add_heading_2("4.3 HTML5 Canvas 2D Vector Drawing Engine (`apps/excel_front/draw/Game.ts`)")
    add_body_p("The `Game.ts` class contains over 780 lines of production-grade vector graphic calculations, interactive event handling, and viewport transformation logic.")

    add_heading_3("1. Viewport Transformation & High-DPI Matrix Scaling")
    add_body_p(
        "To guarantee crisp rendering on high-density Retina and High-DPI screens, `Game.ts` computes the window's `devicePixelRatio` (`dpr = window.devicePixelRatio || 1`). Upon every frame rendering pass, `render()` clears device space (`setTransform(1, 0, 0, 1, 0, 0)`) and sets world coordinate scaling:"
    )
    add_code_block("const s = this.scale * this.dpr;\nctx.setTransform(s, 0, 0, s, this.offsetX * this.dpr, this.offsetY * this.dpr);")

    add_heading_3("2. Coordinate System Mapping (`toWorld` & `screenPoint`)")
    add_body_p(
        "Mouse interaction events occur in screen pixel space (`clientX`, `clientY`). The engine converts screen pixels to world canvas coordinates using inverse translation and scaling:"
    )
    add_code_block("private toWorld(sx: number, sy: number): Point {\n  return {\n    x: (sx - this.offsetX) / this.scale,\n    y: (sy - this.offsetY) / this.scale\n  };\n}")

    add_heading_3("3. Vector Shape Geometry & Wire Message Protocol")
    add_body_p(
        "Every drawn primitive is represented as a discriminated union type `Shape`, extended with styling properties (`id`, `strokeColor`, `strokeWidth`):"
    )
    add_code_block(
        "export type Shape = StyledShape & (\n"
        "  | { type: 'rect'; x: number; y: number; width: number; height: number }\n"
        "  | { type: 'circle'; centerX: number; centerY: number; radius: number }\n"
        "  | { type: 'line'; startX: number; startY: number; endX: number; endY: number }\n"
        "  | { type: 'rhombus'; x: number; y: number; width: number; height: number }\n"
        "  | { type: 'pencil'; points: Point[] }\n"
        ");\n\n"
        "type WireMessage =\n"
        "  | Shape\n"
        "  | { type: 'delete'; ids: string[] }\n"
        "  | { type: 'update'; shape: Shape };"
    )

    add_heading_3("4. Freehand Quadratic Path Smoothing Algorithm")
    add_body_p(
        "Rather than rendering jagged polyline segments for pencil drawings, `Game.ts` computes quadratic Bezier curves passing through segment midpoints (`drawSmoothPath`):"
    )
    add_code_block(
        "ctx.moveTo(points[0].x, points[0].y);\n"
        "for (let i = 1; i < points.length - 1; i++) {\n"
        "  const midX = (points[i].x + points[i + 1].x) / 2;\n"
        "  const midY = (points[i].y + points[i + 1].y) / 2;\n"
        "  ctx.quadraticCurveTo(points[i].x, points[i].y, midX, midY);\n"
        "}\n"
        "ctx.lineTo(points[points.length - 1].x, points[points.length - 1].y);"
    )

    add_heading_3("5. Interactive Hit-Testing & Eraser Distance Math")
    add_body_p(
        "Shape selection and erasing rely on line segment projection mathematics (`distToSegment` and `distanceToShape`). The distance from a cursor point `(px, py)` to segment `(ax, ay)-(bx, by)` is computed via vector projection clamping `t = clamp(((px - ax)*dx + (py - ay)*dy) / lenSq, 0, 1)`."
    )

    add_heading_3("6. Cinematic Dark-Mode Neon Glow Engine")
    add_body_p(
        "In dark canvas mode, every shape automatically gains a dynamic colored halo effect during rendering passes:"
    )
    add_code_block(
        "if (this.theme === 'dark' && !selected) {\n"
        "  ctx.shadowColor = color;\n"
        "  ctx.shadowBlur = 8 + (shape.strokeWidth ?? DEFAULT_WIDTH) * 1.5;\n"
        "}"
    )

    # -------------------------------------------------------------
    # SECTION 5: STEP-BY-STEP INSTALLATION GUIDE
    # -------------------------------------------------------------
    add_heading_1("5. Step-by-Step Installation Guide")

    add_body_p(
        "Follow these exact commands to install dependencies, configure environment variables, and launch the complete monorepo locally."
    )

    add_heading_2("5.1 System Prerequisites")
    add_bullet_p("v20.x or higher (enforced in root `package.json`).", bold_prefix="Node.js: ")
    add_bullet_p("v9.0.0 or higher (`corepack enable pnpm` or `npm i -g pnpm`).", bold_prefix="pnpm: ")
    add_bullet_p("Local or hosted instance (Neon, Supabase, Railway, or Docker).", bold_prefix="PostgreSQL Database: ")

    add_heading_2("5.2 Repository Cloning & Dependency Setup")
    add_code_block(
        "# Step 1: Clone the ExcelDraw repository\n"
        "git clone https://github.com/HarshitJain-hbtu/ExcelDraw.git\n"
        "cd ExcelDraw\n\n"
        "# Step 2: Install monorepo dependencies via pnpm\n"
        "pnpm install"
    )

    add_heading_2("5.3 Environment Configuration (`.env`)")
    add_body_p("Create the required environment files across packages and applications:")

    add_body_p("1. Database Package Environment (`packages/database/.env`):", bold_prefix="")
    add_code_block("DATABASE_URL=\"postgresql://postgres:YOUR_PASSWORD@localhost:5432/exceldraw?schema=public\"")

    add_body_p("2. HTTP Backend & WebSocket Secrets (System Shell or App Environment):", bold_prefix="")
    add_code_block(
        "# PowerShell (Windows)\n"
        "$env:JWT_SECRET=\"your-secure-random-secret-key-32-chars\"\n"
        "$env:PORT=\"4000\"\n\n"
        "# Bash / Linux / macOS\n"
        "export JWT_SECRET=\"your-secure-random-secret-key-32-chars\"\n"
        "export PORT=\"4000\""
    )

    add_heading_2("5.4 Database Schema Push & Client Generation")
    add_code_block(
        "# Navigate to database package and push schema to PostgreSQL\n"
        "cd packages/database\n"
        "pnpm prisma db push\n\n"
        "# Generate type-safe Prisma client\n"
        "pnpm prisma generate\n"
        "cd ../.."
    )

    add_heading_2("5.5 Launching the Development Ecosystem")
    add_code_block(
        "# Start all applications simultaneously using Turborepo\n"
        "pnpm dev"
    )

    add_callout(
        "Running `pnpm dev` executes Turbo's concurrent pipeline, spinning up:\n"
        "• Next.js Frontend (`apps/excel_front`): http://localhost:3000\n"
        "• Express HTTP Backend (`apps/http-backend`): http://localhost:4000\n"
        "• Realtime WebSocket Server (`apps/websockets`): ws://localhost:8080",
        title="LOCAL PORTS OVERVIEW"
    )

    # -------------------------------------------------------------
    # SECTION 6: USAGE GUIDE
    # -------------------------------------------------------------
    add_heading_1("6. Usage Guide")

    add_body_p(
        "This section provides step-by-step instructions for end-users and developers interacting with a running ExcelDraw instance."
    )

    add_heading_2("6.1 End-User Operational Workflow")
    add_bullet_p("Open `http://localhost:3000/auth` in your browser. Switch between Sign In and Sign Up tabs. Provide name, email, and password to authenticate.", bold_prefix="1. User Authentication: ")
    add_bullet_p("Upon login, you are redirected to `/dashboard`. Click the **'+ Create Room'** button, type a unique workspace slug (e.g. `architecture-diagram`), and click Create.", bold_prefix="2. Workspace Dashboard: ")
    add_bullet_p("Click any room card on the dashboard to open the infinite canvas workspace at `/canvas/[roomId]`.", bold_prefix="3. Room Entry: ")
    add_bullet_p("Select drawing tools from the top glassmorphism toolbar or use keyboard shortcuts. Click and drag to create shapes.", bold_prefix="4. Vector Drawing: ")
    add_bullet_p("Copy the URL from the browser bar and send it to teammates. Any authenticated user joining the same room slug will see edits synchronized in real-time.", bold_prefix="5. Multi-User Collaboration: ")

    add_heading_2("6.2 Keyboard Shortcuts Quick Reference")

    kb_table = doc.add_table(rows=1, cols=3)
    kb_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(kb_table)

    kb_headers = ["Shortcut Key", "Action / Tool Mode", "Engine Behavior Description"]
    kb_hdr_cells = kb_table.rows[0].cells
    for i, h in enumerate(kb_headers):
        kb_hdr_cells[i].width = Inches([1.5, 2.0, 3.0][i])
        set_cell_background(kb_hdr_cells[i], HEX_PRIMARY)
        set_cell_margins(kb_hdr_cells[i], top=140, bottom=140, left=120, right=120)
        p = kb_hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h)
        run.font.name = "Arial"
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGB_WHITE

    kb_data = [
        ("V", "Select Tool", "Enables shape selection, hit-testing, and click-and-drag translation."),
        ("P", "Pencil Tool", "Enables freehand smooth quadratic path drawing."),
        ("R", "Rectangle Tool", "Enables box and rectangle vector drawing."),
        ("O / C", "Circle Tool", "Enables center-to-radius circle vector drawing."),
        ("L", "Line Tool", "Enables straight line segment drawing."),
        ("D", "Rhombus Tool", "Enables diamond / rhombus flowchart decision node drawing."),
        ("E", "Eraser Tool", "Enables path proximity erasing with real-time delete broadcast."),
        ("H / Space (Hold)", "Hand / Pan Tool", "Enables smooth infinite workspace panning."),
        ("Delete / Backspace", "Delete Selection", "Removes currently selected shape and broadcasts deletion marker."),
        ("Ctrl / Cmd + Wheel", "Zoom Viewport", "Pinch / scroll zoom centered on cursor (0.1x to 8.0x scale factor).")
    ]

    for idx, row in enumerate(kb_data):
        row_cells = kb_table.add_row().cells
        bg_color = HEX_LIGHT_BG if idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate(row):
            row_cells[i].width = Inches([1.5, 2.0, 3.0][i])
            set_cell_background(row_cells[i], bg_color)
            set_cell_margins(row_cells[i], top=100, bottom=100, left=120, right=120)
            p = row_cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(val)
            run.font.name = "Arial"
            run.font.size = Pt(9.0)
            run.font.color.rgb = RGB_DARK
            if i == 0:
                run.font.bold = True

    doc.add_paragraph()

    # -------------------------------------------------------------
    # SECTION 7: DIRECTORY STRUCTURE
    # -------------------------------------------------------------
    add_heading_1("7. Directory Structure")

    add_body_p(
        "Below is a complete mapped-out tree of the ExcelDraw monorepo workspace explaining the exact purpose and responsibility of every major directory and critical file:"
    )

    tree_str = (
        "ExcelDraw/\n"
        "├── .env.example                # Central environment documentation & variable specs\n"
        "├── .gitignore                  # Git ignore specifications for monorepo build artifacts\n"
        "├── .npmrc                      # PNPM configuration settings\n"
        "├── package.json                # Root Turborepo manifest, scripts, and devDependencies\n"
        "├── pnpm-lock.yaml              # Monorepo lockfile enforcing strict package versions\n"
        "├── pnpm-workspace.yaml         # PNPM workspace definition (`apps/*`, `packages/*`)\n"
        "├── render.yaml                 # Render PaaS deployment specs for HTTP API & WS server\n"
        "├── turbo.json                  # Turborepo task execution pipeline & caching rules\n"
        "├── apps/\n"
        "│   ├── excel_front/            # Next.js 15 (App Router) Frontend Application\n"
        "│   │   ├── app/                # App Router routes and pages\n"
        "│   │   │   ├── auth/           # Authentication pages (signin & signup tabs)\n"
        "│   │   │   ├── canvas/[roomId] # Dynamic canvas route mounting RoomCanvas\n"
        "│   │   │   ├── dashboard/      # User workspace dashboard displaying room cards\n"
        "│   │   │   ├── globals.css     # Tailwind CSS base styles & canvas CSS variables\n"
        "│   │   │   └── layout.tsx      # Root HTML layout with providers & theme wrapper\n"
        "│   │   ├── components/         # React UI component library\n"
        "│   │   │   ├── Canvas.tsx      # Core canvas React wrapper, shortcuts, zoom controls\n"
        "│   │   │   ├── RoomCanvas.tsx  # WebSocket connection manager & auth guard\n"
        "│   │   │   ├── canvas/         # Canvas toolbar (`navbar.tsx` drawing tool selector)\n"
        "│   │   │   ├── cinematic/      # Framer motion & particle background components\n"
        "│   │   │   └── ui/             # Radix UI primitives & custom styled controls\n"
        "│   │   ├── draw/               # HTML5 2D Vector Drawing Engine\n"
        "│   │   │   ├── Game.ts         # 780+ lines canvas engine, hit-testing, zoom, math\n"
        "│   │   │   └── http.ts         # REST API fetch helper for historical chat state\n"
        "│   │   ├── config.ts           # Frontend backend endpoint configuration constants\n"
        "│   │   ├── tailwind.config.ts  # Tailwind CSS theme extension configuration\n"
        "│   │   └── vercel.json         # Vercel deployment configuration\n"
        "│   ├── http-backend/           # Express.js REST API Server\n"
        "│   │   ├── src/\n"
        "│   │   │   ├── index.ts        # Express app initialization, CORS, auth & room routes\n"
        "│   │   │   ├── middleware.ts   # JWT authorization header/cookie extractor middleware\n"
        "│   │   │   └── global.d.ts     # Express Request type extension (`req.userId`)\n"
        "│   │   └── package.json        # Backend scripts & dependency manifest\n"
        "│   └── websockets/             # Real-Time WebSocket Server\n"
        "│       ├── src/\n"
        "│       │   └── index.ts        # Native WS server, token handshake guard, broadcast\n"
        "│       └── package.json        # WebSocket server package manifest\n"
        "└── packages/\n"
        "    ├── backend-common/         # Shared Backend Configuration Package\n"
        "    │   └── src/index.ts        # Exports centralized `JWT_SECRET` constant\n"
        "    ├── common/                 # Shared Data Validation Package\n"
        "    │   └── src/index.ts        # Exports Zod schemas (`CreateUser`, `Signin`, `CreateRoom`)\n"
        "    ├── database/               # Shared Database Access Layer Package\n"
        "    │   ├── prisma/\n"
        "    │   │   └── schema.prisma   # PostgreSQL datasource & Prisma model definitions\n"
        "    │   └── src/db.ts           # Exports singleton `prismaClient` instance\n"
        "    ├── eslint-config/          # Shared ESLint rules package\n"
        "    ├── typescript-config/      # Shared tsconfig base configurations\n"
        "    └── ui/                     # Shared monorepo UI components"
    )

    add_code_block(tree_str)

    # Save document
    output_filename = "ExcelDraw_Technical_Documentation.docx"
    doc.save(output_filename)
    print(f"Successfully generated {output_filename}")

if __name__ == "__main__":
    create_documentation()
